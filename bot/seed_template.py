import json
from bot.database.core import async_session
from bot.database.models import DocumentTemplate, Lawyer, User, Role
from sqlalchemy import select
import asyncio
from docx import Document

def create_example_docx(path: str):
    document = Document()
    document.add_heading('Application Form', 0)

    p = document.add_paragraph('Client Name: ')
    p.add_run('{{ client_name }}').bold = True

    p = document.add_paragraph('Passport: ')
    p.add_run('{{ passport_number }}').bold = True

    p = document.add_paragraph('Address: ')
    p.add_run('{{ address }}').bold = True

    p = document.add_paragraph('Date: ')
    p.add_run('{{ date }}').bold = True

    document.add_paragraph('Request:')
    p = document.add_paragraph()
    p.add_run('{{ request_text }}').italic = True

    document.add_page_break()
    document.save(path)

async def seed():
    create_example_docx("example_template.docx")

    questions = [
        {"key": "client_name", "question": "What is your full name?", "type": "text", "required": True},
        {"key": "passport_number", "question": "What is your passport number? (e.g. AA1234567)", "type": "text", "required": True, "validation": "^[A-Za-z]{2}\\d{7}$"},
        {"key": "address", "question": "What is your address?", "type": "text", "required": True},
        {"key": "date", "question": "What is the date? (DD.MM.YYYY)", "type": "text", "required": True, "validation": "^\\d{2}\\.\\d{2}\\.\\d{4}$"},
        {"key": "request_text", "question": "What is your request?", "type": "text", "required": True}
    ]

    async with async_session() as session:
        # Check if we have a lawyer
        stmt = select(Lawyer).limit(1)
        result = await session.execute(stmt)
        lawyer = result.scalar_one_or_none()

        if not lawyer:
            user = User(telegram_id=1, role=Role.lawyer)
            session.add(user)
            await session.flush()
            lawyer = Lawyer(user_id=user.id, invite_token="test_token")
            session.add(lawyer)
            await session.flush()

        template = DocumentTemplate(
            lawyer_id=lawyer.id,
            name="Example Application",
            file_path="example_template.docx",
            questions=questions
        )
        session.add(template)
        await session.commit()
        print("Example template seeded.")

if __name__ == "__main__":
    asyncio.run(seed())